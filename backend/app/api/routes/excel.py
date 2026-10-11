import csv
import json
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from zipfile import BadZipFile, ZipFile

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.security import audit, permissions, require, scoped
from app.db.session import get_db
from app.models import ImportJob, Movement, Product, Sale, Stock
from app.schemas.sistema import ImportInput, ProductInput
from app.services.catalogo import product_out, save_product, validate_product
from app.services.common import commit, row_out

router = APIRouter(prefix="/excel", tags=["Excel"])
HEADERS = [
    "sku",
    "nombre",
    "descripcion",
    "marca",
    "unidad",
    "tipo_id",
    "categoria_id",
    "producto_padre_id",
    "atributos",
    "precio",
    "costo",
    "stock_minimo",
    "activo",
]
NUMERIC_HEADERS = {
    "precio",
    "costo",
    "stock_minimo",
    "cantidad",
    "saldo",
    "subtotal",
    "impuesto",
    "total",
    "pagado",
    "importe",
    "descuento",
}


def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def workbook(rows, headers):
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(headers)
    for row in rows:
        values = []
        for key in headers:
            value = row.get(key)
            if key in NUMERIC_HEADERS and isinstance(value, str):
                try:
                    value = Decimal(value)
                except InvalidOperation:
                    pass
            values.append(safe_cell(value))
        ws.append(values)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(
            45, max(15, len(str(col[0].value)) + 5)
        )
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def xlsx_response(data, name):
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'},
    )


@router.get("/plantilla")
def template(user=Depends(require("excel.importar"))):
    return xlsx_response(workbook([], HEADERS), "plantilla-productos")


@router.post("/prevalidar")
def preview(
    file: UploadFile = File(...),
    user=Depends(require("excel.importar")),
    db: Session = Depends(get_db),
):
    if "productos.editar" not in permissions(db, user):
        raise HTTPException(403, "Importar requiere gestionar productos")
    content = file.file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "El límite es 5 MB")
    name = (file.filename or "").lower()
    try:
        if name.endswith(".xlsx"):
            with ZipFile(BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(413, "Excel descomprimido demasiado grande")
            wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
            ws = wb.active
            iterator = ws.iter_rows()
            cells = next(iterator, [])
            headers = [str(c.value or "").strip() for c in cells]
            rows = []
            for cells in iterator:
                if any(c.data_type == "f" for c in cells):
                    raise HTTPException(422, "No se admiten fórmulas en la importación")
                row = dict(zip(headers, [c.value for c in cells]))
                rows.append(row)
                if len(rows) > 1000:
                    raise HTTPException(413, "Máximo 1000 filas por importación")
            wb.close()
        elif name.endswith(".csv"):
            reader = csv.DictReader(StringIO(content.decode("utf-8-sig")))
            headers = reader.fieldnames or []
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > 1000:
                    raise HTTPException(413, "Máximo 1000 filas")
        else:
            raise HTTPException(422, "Selecciona un XLSX o CSV UTF-8")
    except (BadZipFile, ValueError, UnicodeError, KeyError, TypeError):
        raise HTTPException(422, "Archivo inválido o dañado") from None
    if (
        not {"sku", "nombre"}.issubset(headers)
        or set(headers) - set(HEADERS)
        or len(headers) != len(set(headers))
    ):
        raise HTTPException(422, "Encabezados incorrectos; usa la plantilla")
    output = []
    errors = []
    seen = set()
    for index, row in enumerate(rows, start=2):
        if all(v is None or v == "" for v in row.values()):
            continue
        try:
            data = {k: v for k, v in row.items() if v is not None and v != ""}
            for k in ("tipo_id", "categoria_id", "producto_padre_id"):
                if k in data:
                    data[k] = int(data[k])
            if "atributos" in data and isinstance(data["atributos"], str):
                data["atributos"] = json.loads(data["atributos"])
            if "activo" in data and isinstance(data["activo"], str):
                values = {
                    "true": True,
                    "1": True,
                    "si": True,
                    "sí": True,
                    "false": False,
                    "0": False,
                    "no": False,
                }
                raw = data["activo"].strip().lower()
                if raw not in values:
                    raise ValueError("activo debe ser true/false, sí/no o 1/0")
                data["activo"] = values[raw]
            p = ProductInput(**data)
            existing = db.scalar(select(Product).where(Product.sku == p.sku))
            validate_product(db, p, existing.id if existing else None)
            if p.sku in seen:
                raise ValueError("SKU repetido en el archivo")
            if p.costo is not None and "costos.ver" not in permissions(db, user):
                raise ValueError("No tienes permiso para importar costos")
            seen.add(p.sku)
            output.append(p.model_dump(mode="json"))
        except (ValidationError, ValueError, HTTPException) as e:
            msg = e.detail if isinstance(e, HTTPException) else str(e)
            errors.append({"fila": index, "error": msg[:500]})
    return {
        "filas": output,
        "errores": errors,
        "validas": len(output),
        "aplicable": bool(output) and not errors,
    }


@router.post("/importar")
def import_products(
    data: ImportInput, user=Depends(require("excel.importar")), db: Session = Depends(get_db)
):
    if "productos.editar" not in permissions(db, user):
        raise HTTPException(403, "Importar requiere gestionar productos")
    if len({p.sku for p in data.filas}) != len(data.filas):
        raise HTTPException(422, "SKU duplicado en el lote")
    for p in data.filas:
        existing = db.scalar(select(Product).where(Product.sku == p.sku).with_for_update())
        save_product(db, p, user, existing)
    job = ImportJob(usuario_id=user.id, filas=len(data.filas))
    db.add(job)
    db.flush()
    audit(db, user, "importar", "productos", job.id, str(job.filas))
    commit(db)
    return {"importacion_id": job.id, "filas": job.filas}


@router.get("/exportar/{tabla}")
def export(tabla: str, user=Depends(require("excel.exportar")), db: Session = Depends(get_db)):
    perms = permissions(db, user)
    mapping = {
        "productos": ("productos.ver", Product),
        "existencias": ("inventario.ver", Stock),
        "movimientos": ("inventario.ver", Movement),
        "ventas": ("ventas.ver", Sale),
    }
    if tabla not in mapping:
        raise HTTPException(404, "Exportación desconocida")
    permission, model = mapping[tabla]
    if permission not in perms:
        raise HTTPException(403, "Sin permiso de lectura para esta tabla")
    q = select(model).order_by(model.id)
    if model != Product:
        q = scoped(q, model.almacen_id, user)
    rows = [
        product_out(db, r, user) if model == Product else row_out(r, ("request_id",))
        for r in db.scalars(q)
    ]
    headers = (
        [c.key for c in Product.__table__.columns if c.key != "costo" or "costos.ver" in perms]
        if model == Product
        else [c.key for c in model.__table__.columns if c.key != "request_id"]
    )
    # Excel cannot represent timezone-aware datetimes directly.
    for row in rows:
        for k, v in row.items():
            if hasattr(v, "isoformat"):
                row[k] = v.isoformat()
    return xlsx_response(workbook(rows, headers), tabla)
