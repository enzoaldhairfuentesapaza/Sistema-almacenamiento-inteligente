from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.db.session import engine
from app.main import app
from app.models import Sale, SaleLine

API = "/api/v1"


def post(c, path, data, status=201):
    r = c.post(API + path, json=data)
    assert r.status_code == status, r.text
    return r.json()


def product(c, sku="P1", **kwargs):
    return post(
        c,
        "/productos",
        dict(sku=sku, nombre="Producto de prueba", precio="10.00", costo="4.00", **kwargs),
    )


def movement(c, id, qty="10", **kwargs):
    return post(
        c,
        "/movimientos",
        dict(
            producto_id=id,
            almacen_id=1,
            cantidad=qty,
            tipo="entrada",
            motivo="Carga de prueba",
            **kwargs,
        ),
    )


def sale(c, id, qty="2", **kwargs):
    return post(
        c,
        "/ventas",
        dict(
            request_id=str(uuid4()),
            almacen_id=1,
            items=[dict(producto_id=id, cantidad=qty)],
            **kwargs,
        ),
    )


def test_full_sale_payment_and_cancellation(client):
    p = product(client)
    movement(client, p["id"])
    s = sale(client, p["id"], tasa_impuesto="0.18")
    assert Decimal(str(s["total"])) == Decimal("23.60")
    assert client.get(API + f"/ventas?q=VE-{s['id']:06d}").json()[0]["id"] == s["id"]
    assert client.get(API + "/ventas?q=²").json() == []
    assert client.get(API + "/ventas?q=" + "9" * 100).json() == []
    assert Decimal(str(client.get(API + "/inventario").json()[0]["cantidad"])) == 8
    pay = post(
        client,
        f"/ventas/{s['id']}/pagos",
        {
            "request_id": str(uuid4()),
            "monto": "23.60",
            "metodo": "yape",
            "referencia": "verificado",
        },
    )
    assert Decimal(str(pay["monto"])) == Decimal("23.60")
    receipt = client.get(API + f"/ventas/{s['id']}/comprobante")
    assert receipt.status_code == 200 and receipt.content.startswith(b"%PDF")
    result = post(client, f"/ventas/{s['id']}/anular", {"motivo": "Devolución total"}, 200)
    assert Decimal(str(result["reembolso_pendiente"])) == Decimal("23.60")
    assert Decimal(str(client.get(API + "/inventario").json()[0]["cantidad"])) == 10
    post(client, f"/ventas/{s['id']}/anular", {"motivo": "Segunda anulación"}, 409)
    assert float(client.get(API + "/reportes/resumen").json()["ventas_total"]) == 0


def test_failed_sale_is_atomic(client):
    a = product(client, "A")
    b = product(client, "B")
    movement(client, a["id"], "5")
    movement(client, b["id"], "1")
    r = client.post(
        API + "/ventas",
        json={
            "request_id": str(uuid4()),
            "almacen_id": 1,
            "items": [
                {"producto_id": a["id"], "cantidad": "2"},
                {"producto_id": b["id"], "cantidad": "2"},
            ],
        },
    )
    assert r.status_code == 409
    stocks = client.get(API + "/inventario").json()
    assert sorted(float(x["cantidad"]) for x in stocks) == [1, 5]
    assert client.get(API + "/ventas").json() == []
    assert len(client.get(API + "/movimientos").json()) == 2


def test_sales_and_payments_are_idempotent(client):
    p = product(client)
    movement(client, p["id"])
    key = str(uuid4())
    body = {
        "request_id": key,
        "almacen_id": 1,
        "items": [{"producto_id": p["id"], "cantidad": "1"}],
    }
    one = post(client, "/ventas", body)
    two = post(client, "/ventas", body)
    assert one["id"] == two["id"]
    payment = {"request_id": str(uuid4()), "monto": "10", "metodo": "efectivo"}
    one_pay = post(client, f"/ventas/{one['id']}/pagos", payment)
    two_pay = post(client, f"/ventas/{one['id']}/pagos", payment)
    assert one_pay["id"] == two_pay["id"]
    assert float(client.get(API + "/inventario").json()[0]["cantidad"]) == 9


def test_transfer_is_balanced_and_rolls_back(client):
    store = post(client, "/almacenes", {"sede_id": 1, "codigo": "B", "nombre": "Segundo"})
    p = product(client)
    movement(client, p["id"], "10")
    body = dict(
        producto_id=p["id"],
        almacen_id=1,
        cantidad="4",
        tipo="transferencia",
        destino_id=store["id"],
        motivo="Mover existencias",
    )
    entries = post(client, "/movimientos", body)
    assert len(entries) == 2
    assert sum(float(s["cantidad"]) for s in client.get(API + "/inventario").json()) == 10
    body["cantidad"] = "20"
    post(client, "/movimientos", body, 409)
    assert sorted(float(s["cantidad"]) for s in client.get(API + "/inventario").json()) == [4, 6]


def test_dynamic_fields_validation_and_variants(client):
    t = post(
        client,
        "/tipos",
        {
            "nombre": "Ropa",
            "campos": [
                {
                    "clave": "talla",
                    "etiqueta": "Talla",
                    "tipo": "seleccion",
                    "obligatorio": True,
                    "opciones": ["S", "M", "L"],
                }
            ],
        },
    )
    post(client, "/productos", {"sku": "INVALID", "nombre": "Falta talla", "tipo_id": t["id"]}, 422)
    post(
        client,
        "/productos",
        {
            "sku": "INVALID",
            "nombre": "Talla inválida",
            "tipo_id": t["id"],
            "atributos": {"talla": "XL"},
        },
        422,
    )
    p = product(client, tipo_id=t["id"], atributos={"talla": "M"})
    v = product(
        client, "P1-L", tipo_id=t["id"], producto_padre_id=p["id"], atributos={"talla": "L"}
    )
    assert v["producto_padre_id"] == p["id"]
    qr = client.get(API + f"/productos/{p['id']}/qr")
    assert qr.status_code == 200 and qr.content.startswith(b"\x89PNG")


def test_restricted_user_scopes_and_cost_redaction(client):
    other = post(client, "/sedes", {"codigo": "LIMA", "nombre": "Lima"})
    store = post(client, "/almacenes", {"sede_id": other["id"], "codigo": "L", "nombre": "Lima"})
    p = product(client)
    movement(client, p["id"])
    post(
        client,
        "/movimientos",
        dict(
            producto_id=p["id"],
            almacen_id=store["id"],
            cantidad="30",
            tipo="entrada",
            motivo="Otro local",
        ),
    )
    role = post(
        client,
        "/roles",
        {
            "nombre": "Consulta",
            "permisos": [
                "empresa.ver",
                "sedes.ver",
                "productos.ver",
                "inventario.ver",
                "excel.exportar",
                "ia.usar",
            ],
        },
    )
    u = post(
        client,
        "/usuarios",
        {
            "nombre": "Lector",
            "email": "lector@example.com",
            "password": "Reader-password-123",
            "roles": [role["id"]],
            "sedes": [1],
        },
    )
    with TestClient(app) as c:
        login = post(
            c, "/auth/login", {"email": u["email"], "password": "Reader-password-123"}, 200
        )
        c.headers["x-csrf-token"] = login["csrf"]
        assert "costo" not in c.get(API + "/productos").json()[0]
        assert len(c.get(API + "/inventario").json()) == 1
        # Search preserves the same site scope and excludes unrelated products.
        assert len(c.get(API + "/inventario?q=P1").json()) == 1
        assert c.get(API + "/inventario?q=missing").json() == []
        assert len(c.get(API + "/movimientos?q=P1").json()) == 1
        assert c.get(API + f"/inventario?almacen_id={store['id']}").status_code == 403
        assert c.get(API + "/usuarios").status_code == 403
        assert c.post(API + "/productos", json={"sku": "X", "nombre": "X"}).status_code == 403
        export = c.get(API + "/excel/exportar/productos")
        wb = load_workbook(BytesIO(export.content))
        assert "costo" not in [cell.value for cell in wb.active[1]]
        scoped_export = c.get(API + "/excel/exportar/existencias")
        wb = load_workbook(BytesIO(scoped_export.content))
        assert wb.active.max_row == 2


def test_session_csrf_password_revocation(client):
    assert client.post(API + "/auth/logout", headers={"x-csrf-token": "bad"}).status_code == 403
    result = post(
        client,
        "/auth/password",
        {"actual": "Safe-test-password-123", "nueva": "New-test-password-123"},
        200,
    )
    assert result["ok"]
    assert client.get(API + "/productos").status_code == 401


def test_duplicate_codes_and_invalid_tax(client):
    post(client, "/sedes", {"codigo": "AQP", "nombre": "Duplicate"}, 409)
    p = product(client)
    movement(client, p["id"])
    post(
        client,
        "/ventas",
        {
            "request_id": str(uuid4()),
            "almacen_id": 1,
            "items": [{"producto_id": p["id"], "cantidad": "-1"}],
        },
        422,
    )
    post(
        client,
        "/ventas",
        {
            "request_id": str(uuid4()),
            "almacen_id": 1,
            "tasa_impuesto": "1.1",
            "items": [{"producto_id": p["id"], "cantidad": "1"}],
        },
        422,
    )


def test_excel_prevalidation_atomic_and_formula_rejection(client):
    content = b"sku,nombre,precio\nA,Producto A,10\nB,Producto B,15\n"
    r = client.post(API + "/excel/prevalidar", files={"file": ("items.csv", content, "text/csv")})
    assert r.status_code == 200, r.text
    preview = r.json()
    assert preview["aplicable"]
    invalid_state = client.post(
        API + "/excel/prevalidar",
        files={"file": ("state.csv", b"sku,nombre,activo\nZ,Producto Z,truu\n", "text/csv")},
    ).json()
    assert not invalid_state["aplicable"]
    assert "activo" in invalid_state["errores"][0]["error"]
    post(client, "/excel/importar", {"filas": preview["filas"]}, 200)
    export = client.get(API + "/excel/exportar/productos")
    sheet = load_workbook(BytesIO(export.content)).active
    price_column = [cell.value for cell in sheet[1]].index("precio") + 1
    assert sheet.cell(2, price_column).data_type == "n"
    assert sheet.cell(2, price_column).value == 10
    post(
        client,
        "/excel/importar",
        {"filas": [{"sku": "C", "nombre": "C"}, {"sku": "D", "nombre": "D", "categoria_id": 999}]},
        422,
    )
    assert len(client.get(API + "/productos").json()) == 2
    wb = Workbook()
    ws = wb.active
    ws.append(["sku", "nombre"])
    ws.append(["XX", "=1+2"])
    out = BytesIO()
    wb.save(out)
    assert (
        client.post(
            API + "/excel/prevalidar", files={"file": ("formula.xlsx", out.getvalue())}
        ).status_code
        == 422
    )


def test_purchase_reception(client):
    p = product(client)
    supplier = post(client, "/contactos", {"tipo": "proveedor", "nombre": "Distribuidor"})
    body = {
        "request_id": str(uuid4()),
        "almacen_id": 1,
        "proveedor_id": supplier["id"],
        "referencia": "OC-1",
        "items": [{"producto_id": p["id"], "cantidad": "3", "costo": "5"}],
    }
    first = post(client, "/compras", body)
    second = post(client, "/compras", body)
    assert first["id"] == second["id"]
    assert float(client.get(API + "/inventario").json()[0]["cantidad"]) == 3
    detail = client.get(API + f"/compras/{first['id']}").json()
    assert detail["proveedor"] == "Distribuidor"
    assert Decimal(detail["items"][0]["costo"]) == Decimal("5")
    assert Decimal(detail["items"][0]["importe"]) == Decimal("15")
    company = client.get(API + "/empresa").json()
    company.pop("id")
    company["moneda"] = "USD"
    assert client.put(API + "/empresa", json=company).status_code == 409


def test_forecast_needs_history_and_evaluates_temporally(client):
    assert client.get(API + "/pronosticos").status_code == 422
    p = product(client)
    with Session(engine) as db:
        for i in range(70):
            dt = datetime.now(timezone.utc) - timedelta(days=70 - i)
            amount = Decimal(10 + i % 7)
            s = Sale(
                request_id=str(uuid4()),
                almacen_id=1,
                usuario_id=1,
                subtotal=amount,
                impuesto=0,
                total=amount,
                pagado=0,
                moneda="PEN",
                fecha=dt,
            )
            db.add(s)
            db.flush()
            db.add(
                SaleLine(
                    venta_id=s.id,
                    producto_id=p["id"],
                    sku="P1",
                    nombre="Test",
                    cantidad=1,
                    precio=amount,
                    costo=4,
                    descuento=0,
                    importe=amount,
                )
            )
        db.commit()
    r = client.get(API + "/pronosticos?horizonte=7")
    assert r.status_code == 200, r.text
    result = r.json()
    assert len(result["predicciones"]) == 7 and len(result["metricas"]) == 2
    assert result["predicciones"][0]["fecha"] <= datetime.now(timezone.utc).date().isoformat()
    assert all(p["inferior"] >= 0 for p in result["predicciones"])


def test_assistant_not_configured_is_explicit(client):
    result = post(client, "/asistente", {"pregunta": "¿Cuánto stock hay?"}, 503)
    assert "configurar" in result["detail"]


def test_inactive_store_and_stock_shortage(client):
    p = product(client)
    movement(client, p["id"], "2")
    client.patch(API + "/almacenes/1/estado", json={"activo": False})
    post(
        client,
        "/ventas",
        {
            "request_id": str(uuid4()),
            "almacen_id": 1,
            "items": [{"producto_id": p["id"], "cantidad": "1"}],
        },
        409,
    )
    assert float(client.get(API + "/inventario").json()[0]["cantidad"]) == 2


def test_password_preserves_quotes_and_spaces(client):
    secret = "  Aa\"bb'cc-password  "
    post(
        client,
        "/usuarios",
        {
            "nombre": "Special",
            "email": "special@example.com",
            "password": secret,
            "roles": [],
            "sedes": [],
        },
    )
    with TestClient(app) as c:
        post(c, "/auth/login", {"email": "special@example.com", "password": secret}, 200)
        post(c, "/auth/login", {"email": "special@example.com", "password": secret.strip()}, 401)
